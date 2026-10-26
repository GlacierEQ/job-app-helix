"""theme.py — shared design system for the resume-forge document family.

Resumes, cover letters, and future packet documents share one visual identity:
glacier-teal accent, Calibri, single-column ATS-safe layout. Change the theme
here and every document follows.
"""

from __future__ import annotations

import re
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ACCENT = "0E5F6B"      # deep glacier teal — name, section heads, rules
INK = "1A1A1A"         # body text
MUTED = "595959"       # secondary text
RULE_LIGHT = "B9CFD4"  # hairline rules

BODY_FONT = "Calibri"
NAME_FONT = "Calibri Light"
BODY_SIZE = 10.5
SMALL_SIZE = 9.5

PAGE_W = 8.5
MARGIN_SIDE = 0.65
MARGIN_TB = 0.55
RIGHT_TAB = PAGE_W - 2 * MARGIN_SIDE  # tab stop position for right-aligned dates

BOLD_RE = re.compile(r"(\*\*.+?\*\*)")


def setup_page(doc):
    section = doc.sections[0]
    section.page_width = Inches(PAGE_W)
    section.page_height = Inches(11)
    section.left_margin = Inches(MARGIN_SIDE)
    section.right_margin = Inches(MARGIN_SIDE)
    section.top_margin = Inches(MARGIN_TB)
    section.bottom_margin = Inches(MARGIN_TB)
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_SIZE)
    normal.font.color.rgb = RGBColor.from_string(INK)


def _set_bottom_border(paragraph, color: str, size: str = "8", space: str = "3"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def _letterspace(run, twentieths: int = 60):
    rPr = run._r.get_or_add_rPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), str(twentieths))
    rPr.append(spacing)


def add_runs(paragraph, text: str, size: float = BODY_SIZE,
             color: str = INK, bold: bool = False,
             italic: bool = False, font: str = BODY_FONT):
    """Add runs to a paragraph, honoring **bold** inline markup."""
    for part in BOLD_RE.split(text):
        if not part:
            continue
        is_bold = bold
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part = part[2:-2]
            is_bold = True
        run = paragraph.add_run(part)
        run.font.name = font
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(color)
        run.bold = is_bold
        run.italic = italic
    return paragraph


def _para(doc, size=BODY_SIZE, space_before=0, space_after=4,
          alignment=WD_ALIGN_PARAGRAPH.LEFT, left_indent=None,
          hanging=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.07
    pf.alignment = alignment
    if left_indent is not None:
        pf.left_indent = Inches(left_indent)
    if hanging is not None:
        pf.first_line_indent = Inches(-hanging)
    return p


def add_header(doc, data):
    """Centered name + contact line with accent rule. data: {name, contact[]}."""
    name = _para(doc, space_after=1, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    run = name.add_run(data["name"])
    run.font.name = NAME_FONT
    run.font.size = Pt(25)
    run.font.color.rgb = RGBColor.from_string(ACCENT)
    run.bold = True

    contact = _para(doc, space_after=6, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_runs(contact, "  ·  ".join(data["contact"]),
             size=SMALL_SIZE, color=MUTED)
    _set_bottom_border(contact, ACCENT, size="12", space="6")


def add_section_head(doc, title: str):
    p = _para(doc, space_before=10, space_after=5)
    run = p.add_run(title.upper())
    run.font.name = BODY_FONT
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(ACCENT)
    run.bold = True
    run.font.small_caps = True
    _letterspace(run, 70)
    _set_bottom_border(p, RULE_LIGHT, size="6", space="2")
