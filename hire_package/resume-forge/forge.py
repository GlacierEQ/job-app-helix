#!/usr/bin/env python3
"""forge.py — resume renderer. Design system lives in theme.py; content in content/*.yaml.

Usage:
  python3 forge.py --variant master        # -> out/Casey_Barton_Resume_Master.docx
  python3 forge.py --variant scale-fde      # -> out/Casey_Barton_Resume_ScaleAI.docx
  python3 forge.py --variant master --pdf   # also render PDF via LibreOffice
"""

from __future__ import annotations

import argparse
import subprocess
import sys
from pathlib import Path

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Inches

from theme import (ACCENT, BODY_SIZE, INK, MUTED, RIGHT_TAB, SMALL_SIZE,
                   _para, add_header, add_runs, add_section_head, setup_page)

HERE = Path(__file__).resolve().parent
CONTENT_DIR = HERE / "content"
OUT_DIR = HERE / "out"


def add_profile(doc, lines):
    for line in lines:
        p = _para(doc, space_after=3)
        add_runs(p, line)


def add_competencies(doc, groups):
    for g in groups:
        p = _para(doc, space_after=2)
        add_runs(p, f"{g['category']}: ", bold=True, size=BODY_SIZE)
        add_runs(p, g["items"], size=BODY_SIZE)


def add_experience(doc, roles):
    for role in roles:
        p = _para(doc, space_before=6, space_after=0)
        p.paragraph_format.tab_stops.add_tab_stop(
            Inches(RIGHT_TAB), WD_TAB_ALIGNMENT.RIGHT)
        add_runs(p, role["title"], bold=True)
        add_runs(p, "\t" + role["dates"], color=MUTED, size=SMALL_SIZE)
        q = _para(doc, space_after=3)
        add_runs(q, f"{role['org']} — {role['location']}",
                 italic=True, color=MUTED, size=SMALL_SIZE)
        for b in role["bullets"]:
            bp = _para(doc, space_after=2, left_indent=0.28, hanging=0.16)
            add_runs(bp, "•  " + b)


def add_education(doc, schools):
    for s in schools:
        p = _para(doc, space_before=5, space_after=0)
        p.paragraph_format.tab_stops.add_tab_stop(
            Inches(RIGHT_TAB), WD_TAB_ALIGNMENT.RIGHT)
        add_runs(p, s["school"], bold=True)
        if s.get("dates"):
            add_runs(p, "\t" + s["dates"], color=MUTED, size=SMALL_SIZE)
        q = _para(doc, space_after=2)
        add_runs(q, s["detail"], italic=True, color=MUTED, size=SMALL_SIZE)


def add_proof(doc, items):
    for item in items:
        p = _para(doc, space_before=6, space_after=1)
        add_runs(p, item["title"], bold=True)
        for b in item["bullets"]:
            bp = _para(doc, space_after=2, left_indent=0.28, hanging=0.16)
            add_runs(bp, "•  " + b)


def add_skills(doc, groups):
    for g in groups:
        p = _para(doc, space_after=2)
        add_runs(p, f"{g['category']}: ", bold=True, size=BODY_SIZE)
        add_runs(p, g["items"], size=BODY_SIZE)


def add_boundary_note(doc, text):
    p = _para(doc, space_before=10, space_after=0)
    add_runs(p, text, size=8.5, color=MUTED, italic=True)


VARIANT_FILES = {
    "master": "Casey_Barton_Resume_Master.docx",
    "scale-fde": "Casey_Barton_Resume_ScaleAI.docx",
    "databricks-fde": "Casey_Barton_Resume_Databricks_FDE.docx",
    "scale-fde-ps": "Casey_Barton_Resume_Scale_FDE_PublicSector.docx",
    "clickup-maf-senior": "Casey_Barton_Resume_ClickUp_MAF_Senior.docx",
    "harvey-agents-staff": "Casey_Barton_Resume_Harvey_Agents_Staff.docx",
    "cresta-fde-ai-agent": "Casey_Barton_Resume_Cresta_FDE_AIAgent.docx",
    "akkio-solutions-architect": "Casey_Barton_Resume_Akkio_SolutionsArchitect.docx",
    "azx-fde-ai": "Casey_Barton_Resume_AZX_SeniorFDE.docx",
    "doximity-staff-ai": "Casey_Barton_Resume_Doximity_StaffAppliedAI.docx",
    "provectus-fde": "Casey_Barton_Resume_Provectus_FDE.docx",
    "climb-fde": "Casey_Barton_Resume_Climb_FDE.docx",
    "mastra-ai-engineer": "Casey_Barton_Resume_Mastra_AIEngineer.docx",
    "stickermule-si-agent": "Casey_Barton_Resume_StickerMule_SIAgent.docx",
    "zoominfo-principal-applied-ai": "Casey_Barton_Resume_ZoomInfo_PrincipalAppliedAI.docx",
    "coinbase-senior-ml-cx": "Casey_Barton_Resume_Coinbase_SeniorML_CX.docx",
    "gc-ai-applied-ai": "Casey_Barton_Resume_GCAI_AppliedAI.docx",
    "topsail-applied-ai": "Casey_Barton_Resume_Topsail_AppliedAI.docx",
    "elevenlabs-creative-agents": "Casey_Barton_Resume_ElevenLabs_CreativeAgents.docx",
    "wrapbook-agentic-systems": "Casey_Barton_Resume_Wrapbook_AgenticSystems.docx",
    "langchain-deployed-architect": "Casey_Barton_Resume_LangChain_DeployedArchitect.docx",
    "cadence-senior-ai-engineer": "Casey_Barton_Resume_Cadence_SeniorAIEngineer.docx",
    "amgen-senior-fde": "Casey_Barton_Resume_Amgen_SeniorFDE.docx",
    "jetbrains-fde": "Casey_Barton_Resume_JetBrains_FDE.docx",
    "reddit-fde-senior": "Casey_Barton_Resume_Reddit_SeniorFDE.docx",
    "reddit-fde-lead": "Casey_Barton_Resume_Reddit_LeadFDE.docx",
    "reddit-agent-opt-eval": "Casey_Barton_Resume_Reddit_AgentOptEval.docx",
    "reddit-agentic-ml": "Casey_Barton_Resume_Reddit_AgenticML_IC5.docx",
    "emax-senior-ai": "Casey_Barton_Resume_eMAX_SeniorAI.docx",
    "pearson-ai-scientist": "Casey_Barton_Resume_Pearson_AIScientist.docx",
    "afresh-staff-fde": "Casey_Barton_Resume_Afresh_StaffFDE.docx",
    "cribl-staff-agentic-tools": "Casey_Barton_Resume_Cribl_StaffAgenticTools.docx",
    "bmc-principal-ai": "Casey_Barton_Resume_BMC_PrincipalAI.docx",
    "future-applied-ai": "Casey_Barton_Resume_Future_AppliedAI.docx",
    "gitlab-staff-fde": "Casey_Barton_Resume_GitLab_StaffFDE.docx",
    "fieldwire-aiml-engineer": "Casey_Barton_Resume_Fieldwire_AIMLEngineer.docx",
    "stripe-solutions-engineer-ai-build": "Casey_Barton_Resume_Stripe_SolutionsEngineer_AIBuild.docx",
    "higharc-fde": "Casey_Barton_Resume_Higharc_FDE.docx",
    "nebius-customer-ai-engineer": "Casey_Barton_Resume_Nebius_CustomerAIEngineer.docx",
    "assemblyai-agent-experience": "Casey_Barton_Resume_AssemblyAI_AgentExperience.docx",
    "cohere-fde-agentic": "Casey_Barton_Resume_Cohere_FDE_Agentic.docx",
}


def build(variant: str) -> Path:
    data = yaml.safe_load((CONTENT_DIR / f"{variant}.yaml").read_text())
    doc = Document()
    setup_page(doc)

    add_header(doc, data)
    add_section_head(doc, "Profile")
    add_profile(doc, data["profile"])

    add_section_head(doc, "Core Competencies")
    add_competencies(doc, data["competencies"])

    add_section_head(doc, "Experience")
    add_experience(doc, data["experience"])

    if data.get("education"):
        add_section_head(doc, "Education")
        add_education(doc, data["education"])

    add_section_head(doc, "Selected Proof")
    add_proof(doc, data["proof"])

    add_section_head(doc, "Skills")
    add_skills(doc, data["skills"])

    if data.get("boundary_note"):
        add_boundary_note(doc, data["boundary_note"])

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / VARIANT_FILES[variant]
    doc.save(out)
    return out


def to_pdf(docx_path: Path) -> Path:
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf",
         "--outdir", str(OUT_DIR), str(docx_path)],
        check=True, capture_output=True, text=True)
    return docx_path.with_suffix(".pdf")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--variant", choices=VARIANT_FILES, required=True)
    ap.add_argument("--pdf", action="store_true")
    args = ap.parse_args()

    out = build(args.variant)
    print(f"DOCX: {out}")
    if args.pdf:
        pdf = to_pdf(out)
        print(f"PDF:  {pdf}")


if __name__ == "__main__":
    sys.exit(main())
