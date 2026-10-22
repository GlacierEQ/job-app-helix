"""Build three DRAFT resume packs for the 2026-10-21 00:42 sweep shortlist.
Style replicates packs-2026-10-20/cohere-fde-agentic/ exactly.
Estate numbers: kit-allowlisted 65 MCP / 989-test / 62-62; 1,398 repos live-verified 2026-10-20.
Honesty: three years never rounded; no weight fine-tuning; no invented depth.
"""
from docx import Document
from docx.shared import Pt, RGBColor

TEAL = RGBColor(0x0E, 0x5F, 0x6B)
GRAY = RGBColor(0x59, 0x59, 0x59)
INK = RGBColor(0x1A, 0x1A, 0x1A)
CONTACT = ("Honolulu, HI  ·  (808) 936-5654  ·  glacier.equilibrium@gmail.com  ·  "
           "github.com/GlacierEQ  ·  casey-barton-glaciereq.vercel.app  ·  linkedin.com/in/caseybartonai")
LS = 1.0708333333333333

def base_doc():
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"; st.font.size = Pt(10.5); st.font.color.rgb = INK
    sec = d.sections[0]
    sec.left_margin = Pt(46.8); sec.right_margin = Pt(46.8)
    sec.top_margin = Pt(39.6); sec.bottom_margin = Pt(39.6)
    return d

def name_p(d, text):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = LS
    r = p.add_run(text); r.bold = True; r.font.size = Pt(25); r.font.color.rgb = TEAL
    return p

def contact_p(d):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = LS
    r = p.add_run(CONTACT); r.font.size = Pt(9.5); r.font.color.rgb = GRAY
    return p

def head(d, text):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(5); p.paragraph_format.line_spacing = LS
    r = p.add_run(text); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = TEAL
    return p

def para(d, runs, before=0, after=2):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = LS
    for text, bold in runs:
        r = p.add_run(text); r.bold = bold; r.font.size = Pt(10.5); r.font.color.rgb = INK
    return p

def comp_line(d, label, items):
    return para(d, [(label + ": ", True), (items, False)])

def bullet(d, runs):
    return para(d, [("•  ", False)] + runs, after=2)

def exp_header(d, title, dates, org, loc):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = LS
    r1 = p.add_run(title + "\t"); r1.bold = True
    r2 = p.add_run(dates); r2.bold = True
    for r in (r1, r2):
        r.font.size = Pt(10.5); r.font.color.rgb = INK
    para(d, [(org + " — " + loc, False)], before=0, after=2)

def edu(d, rows):
    for school, dates, degree in rows:
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = LS
        r1 = p.add_run(school + "\t"); r1.bold = True; r1.font.size = Pt(10.5); r1.font.color.rgb = INK
        r2 = p.add_run(dates); r2.bold = True; r2.font.size = Pt(10.5); r2.font.color.rgb = INK
        para(d, [(degree, False)], before=0, after=2)

EDU_ROWS = [
    ("University of Hawaii at Manoa", "2016", "Bachelor of Science"),
    ("AWS Cloud Institute", "2025", "Cloud Computing program"),
    ("Colorado State University", "", "Undergraduate coursework"),
]

EVIDENCE_BOUNDARY = ("Evidence boundary: ", True)

# ==================== PERPLEXITY MODEL BEHAVIOR ====================
def build_pplx_behavior(path):
    d = base_doc()
    name_p(d, "Casey Barton"); contact_p(d)
    head(d, "PROFILE")
    para(d, [("Applied AI engineer who shapes how models behave in production: context and prompt engineering, "
              "tool/skill/memory design, and self-improvement loops that steer the prompt and improve tool use — "
              "built the experiment-driven way, with failure modes found by well-designed research, not vibes. "
              "Three years as a founder running an evidence-grade agent control plane: I identify inconsistencies and "
              "failure modes in model outputs (stuck loops, bad retrievals, tool misuse, context blowouts), then "
              "design the evals and benchmarks that prove the fix. Written rubrics behind every automated quality "
              "decision; opinions defended by measured numbers.", False)])
    para(d, [("The loop is real: ", False), ("65", True),
              (" authored MCP connectors — governed tool/skill connections with typed contracts — exercised daily by "
              "agents whose behavior I tune, measure, and re-tune. No weight fine-tuning claimed; the behavior-shaping "
              "surface is eval-driven prompt, context, tool, and memory design.", False)])
    head(d, "CORE COMPETENCIES")
    comp_line(d, "Model behavior & context engineering",
              "prompt engineering · context engineering · tool/skill/memory design · self-improvement loops · "
              "behavioral steering · response-style shaping")
    comp_line(d, "Evaluation & research",
              "eval harnesses · golden datasets · behavioral assertions · LLM-as-a-judge with written rubrics · "
              "failure-mode research · inconsistency analysis · benchmarks")
    comp_line(d, "Agentic AI",
              "domain-specific agents · tool/function calling · MCP · structured outputs · RAG · guardrails · "
              "fail-closed recovery · human-in-the-loop gates")
    comp_line(d, "Production",
              "Python (primary) · TypeScript · SQL · Bash · REST APIs · JSON Schema / Pydantic-style validation · "
              "GitHub Actions (CI/CD) · Postgres · latency & token-cost instrumentation")
    head(d, "EXPERIENCE")
    exp_header(d, "Founder / Applied AI Systems Architect", "2023 – Present", "GlacierEQ", "Honolulu, HI")
    bullet(d, [("Shape model behavior across products and use cases: design, test, and optimize the prompts, tools, "
                "skills, and memory that determine response quality — context-engineering strategies tuned per domain "
                "and model, with the tradeoffs measured.", False)])
    bullet(d, [("Build self-improvement loops: instrument agent runs, find where behavior degrades, steer the prompt "
                "and improve tool/skill use and memory recall — the loop compounds instead of resetting every "
                "session.", False)])
    bullet(d, [("Run failure-mode research on production outputs: stuck loops, bad retrievals, tool misuse, context "
                "blowouts, style inconsistencies — identified through designed experiments, then fixed with "
                "behavioral assertions and eval gates that prove the fix held.", False)])
    bullet(d, [("Design the evals that decide what ships: golden datasets, behavioral assertions, LLM-as-a-judge with "
                "written rubrics — quality decisions declared up front, measured, and auditable. A ", False),
               ("989-test", True), (" suite and ", False), ("62/62", True),
               (" verification checks green back the posture across an ", False), ("1398-repository", True),
               (" estate.", False)])
    bullet(d, [("Connect behavior to the real world through ", False), ("65", True),
               (" authored MCP connectors — governed tool/API/data-source connections with typed contracts and "
                "structured validation on every boundary.", False)])
    bullet(d, [("Share the intuition: write the context-engineering and prompt-design practices down so every "
                "engagement inherits them — playbooks, not tribal knowledge.", False)])
    head(d, "EDUCATION"); edu(d, EDU_ROWS)
    head(d, "SELECTED PROOF")
    para(d, [("Behavioral eval infrastructure — quality declared, then measured", True)])
    bullet(d, [("LLM-as-a-judge with written rubrics, golden datasets, behavioral assertions: the exact machinery "
                "behind 'identify inconsistencies and failure modes through well-designed research projects.'", False)])
    para(d, [("Failure posture — designed for production, not the demo", True)])
    bullet(d, [("Stuck loops, bad retrievals, tool misuse, context blowouts: instrumented in production, caught by "
                "replay loops and fail-closed gates, recovery restricted to evidence declared before "
                "reassessment.", False)])
    head(d, "SKILLS")
    comp_line(d, "Languages", "Python · TypeScript · SQL · Bash · JavaScript / Node.js")
    comp_line(d, "Behavior & eval",
              "prompt/context engineering · tool/skill/memory design · self-improvement loops · eval harnesses · "
              "LLM-as-a-judge · golden datasets · behavioral assertions · trace intelligence")
    comp_line(d, "Production",
              "REST APIs · MCP · Postgres · CI/CD (GitHub Actions) · JSON Schema validation · observability · "
              "latency & token-cost instrumentation · retry logic")
    para(d, [EVIDENCE_BOUNDARY,
             ("portfolio repositories are independent work. Three years stated as three years — inside the posting's "
              "2–10+ year band, honestly. No weight fine-tuning claimed — behavior shaping is eval-driven prompt, "
              "context, tool, and memory design. No Perplexity-internal claims; transferable behavior/eval engineering "
              "only. No Go / GCP / GraphQL / K8s / Terraform / DS-statistics depth claimed.", False)])
    d.save(path)

# ==================== PERPLEXITY ANSWER QUALITY & EVALS ====================
def build_pplx_evals(path):
    d = base_doc()
    name_p(d, "Casey Barton"); contact_p(d)
    head(d, "PROFILE")
    para(d, [("Evaluation engineer for answer quality: I build the measurement systems that tell an AI product whether "
              "its answers are actually good — golden datasets, LLM-as-a-judge with written rubrics, behavioral "
              "assertions, and trace intelligence that turns production runs into the next eval iteration. Three years "
              "as a founder running an evidence-grade agent control plane where nothing ships until the evals say it "
              "should: fail-closed GO / NO-GO decisions, written rubrics behind every automated quality call. "
              "Answer quality is the product metric I optimize — measured, not vibed.", False)])
    para(d, [("The systems are real: ", False), ("65", True),
              (" authored MCP connectors feeding agents whose answer quality I measure daily; a ", False),
              ("989-test", True), (" suite and ", False), ("62/62", True),
              (" verification checks green across a ", False), ("1398-repository", True),
              (" estate. Production systems shipped and operated — the 'shipping and operating production systems' "
               "bar, honestly stated at founder/operator scale.", False)])
    head(d, "CORE COMPETENCIES")
    comp_line(d, "Evaluation & answer quality",
              "eval harnesses · LLM-as-a-judge with written rubrics · golden datasets · behavioral assertions · "
              "answer-quality metrics · regression gates · trace intelligence")
    comp_line(d, "Production ML systems",
              "shipping and operating production systems · observability · latency & token-cost instrumentation · "
              "CI/CD (GitHub Actions) · incident-shaped thinking")
    comp_line(d, "Agentic AI",
              "domain-specific agents · tool/function calling · MCP · structured outputs · RAG · guardrails · "
              "fail-closed recovery")
    comp_line(d, "Production",
              "Python (primary) · TypeScript · SQL · Bash · REST APIs · JSON Schema / Pydantic-style validation · "
              "Postgres · retry logic")
    head(d, "EXPERIENCE")
    exp_header(d, "Founder / Applied AI Systems Architect", "2023 – Present", "GlacierEQ", "Honolulu, HI")
    bullet(d, [("Own answer quality end to end: define what 'good' means in written rubrics, build the judges and "
                "datasets that measure it, gate releases on the numbers — the eval decides what ships.", False)])
    bullet(d, [("Ship and operate production AI systems: agents in live operator workflows, monitored, iterated, and "
                "held to quality bars — three years of production ownership at founder/operator scale.", False)])
    bullet(d, [("Turn production into evals: trace intelligence over real runs feeds golden-dataset growth and rubric "
                "refinement — the measurement system improves with every deployment.", False)])
    bullet(d, [("Instrument the failure modes that degrade answers: stuck loops, bad retrievals, tool misuse, "
                "context blowouts — caught by behavioral assertions and fail-closed gates before users see "
                "them.", False)])
    bullet(d, [("Connect measurement to the real world through ", False), ("65", True),
               (" authored MCP connectors — governed tool/API/data-source connections with typed contracts, "
                "structured validation, and audit receipts.", False)])
    head(d, "EDUCATION"); edu(d, EDU_ROWS)
    head(d, "SELECTED PROOF")
    para(d, [("Evaluation-gated quality — the eval decides what ships", True)])
    bullet(d, [("Fail-closed GO / NO-GO campaign engine: immutable evidence models, per-run verification receipts — "
                "the discipline of declared success criteria applied to every production engagement.", False)])
    para(d, [("Answer-quality machinery — rubrics, judges, datasets", True)])
    bullet(d, [("LLM-as-a-judge with written rubrics, golden datasets, behavioral assertions: directly shapes how "
                "answer quality is measured and improved — the posting's core sentence, as built systems.", False)])
    head(d, "SKILLS")
    comp_line(d, "Languages", "Python · TypeScript · SQL · Bash · JavaScript / Node.js")
    comp_line(d, "Eval & quality",
              "eval harnesses · LLM-as-a-judge · written rubrics · golden datasets · behavioral assertions · "
              "regression gates · trace intelligence · failure analysis")
    comp_line(d, "Production",
              "REST APIs · MCP · Postgres · CI/CD (GitHub Actions) · JSON Schema validation · observability · "
              "latency & token-cost instrumentation")
    para(d, [EVIDENCE_BOUNDARY,
             ("portfolio repositories are independent work. Three years stated as three years — against the posting's "
              "4+ year bar, honestly stated; the eval/production-ownership depth is the case. No weight fine-tuning "
              "claimed. No Perplexity-internal claims; transferable eval/quality engineering only. No Go / GCP / "
              "GraphQL / K8s / Terraform / DS-statistics depth claimed.", False)])
    d.save(path)

# ==================== VAPI AGENTIC DEVELOPER EXPERIENCE ====================
def build_vapi_dx(path):
    d = base_doc()
    name_p(d, "Casey Barton"); contact_p(d)
    head(d, "PROFILE")
    para(d, [("Builder of the agentic development lifecycle: I turn fragmented agent pieces — sandboxes, CI, "
              "verification, background agents — into one reliable system where scoped tasks become review-ready work "
              "backed by real evidence. Three years as a founder running an evidence-grade agent control plane: "
              "automated quality gates decide what ships, telemetry and trace intelligence prove it, and the harness "
              "compounds across every engagement. Developer tooling that makes agents trustworthy is the product I "
              "build.", False)])
    para(d, [("The harness is real: ", False), ("65", True),
              (" authored MCP connectors — the governed tool/API surface agents build against — plus eval harnesses "
              "with written rubrics (LLM-as-a-judge), golden datasets, and behavioral assertions as the verification "
              "layer. CI/CD (GitHub Actions), deterministic regression gates, and fail-closed recovery across a ",
               False), ("1398-repository", True), (" estate with a ", False), ("989-test", True),
              (" suite and ", False), ("62/62", True), (" verification checks green.", False)])
    head(d, "CORE COMPETENCIES")
    comp_line(d, "Agentic developer experience",
              "agent harnesses · sandbox environments · developer tooling · validation infrastructure · automated "
              "quality gates · background agents · task-to-production instrumentation")
    comp_line(d, "Verification & quality",
              "eval harnesses · LLM-as-a-judge with written rubrics · golden datasets · behavioral assertions · "
              "trace intelligence · telemetry · fail-closed recovery")
    comp_line(d, "Platform engineering",
              "API design · CI/CD (GitHub Actions) · monorepo-scale workflows · internal platforms · developer "
              "productivity instrumentation")
    comp_line(d, "Production",
              "Python (primary) · TypeScript · SQL · Bash · REST APIs · JSON Schema / Pydantic-style validation · "
              "Postgres · observability · retry logic")
    head(d, "EXPERIENCE")
    exp_header(d, "Founder / Applied AI Systems Architect", "2023 – Present", "GlacierEQ", "Honolulu, HI")
    bullet(d, [("Own the agent development lifecycle: map the path from task to production, instrument the "
                "constraints on agent effectiveness, and ship the sandbox, tooling, and validation improvements that "
                "remove them.", False)])
    bullet(d, [("Build the verification layer agents ship against: automated quality gates with written rubrics, "
                "behavioral assertions, and trace-backed evidence — scoped tasks become review-ready work backed by "
                "real evidence, not assertions.", False)])
    bullet(d, [("Give background agents what they need: context, tools, sandbox, telemetry, and the gates that turn "
                "their output into shippable work — the internal software-factory motion, built and operated.", False)])
    bullet(d, [("Measure agent effectiveness and iterate: latency, quality, and cost instrumentation over every run — "
                "the harness improves because the numbers say where it hurts.", False)])
    bullet(d, [("Connect it all through ", False), ("65", True),
               (" authored MCP connectors — governed integrations with typed contracts, the tool surface the harness "
                "validates against.", False)])
    head(d, "EDUCATION"); edu(d, EDU_ROWS)
    head(d, "SELECTED PROOF")
    para(d, [("Agentic development lifecycle — task to production, instrumented", True)])
    bullet(d, [("Job-App Helix campaign engine: per-run verification receipts, immutable evidence models, fail-closed "
                "GO / NO-GO — the pipeline that turns agent work into shippable, evidenced output.", False)])
    para(d, [("Verification infrastructure — gates, judges, traces", True)])
    bullet(d, [("Eval harnesses with written rubrics, golden datasets, behavioral assertions, trace intelligence: the "
                "validation infrastructure behind every agent engagement I run.", False)])
    head(d, "SKILLS")
    comp_line(d, "Languages", "Python · TypeScript · SQL · Bash · JavaScript / Node.js")
    comp_line(d, "Agentic DX",
              "agent harnesses · sandboxes · developer tooling · validation infra · quality gates · background-agent "
              "orchestration · telemetry · CI/CD")
    comp_line(d, "Production",
              "REST APIs · MCP · Postgres · GitHub Actions · JSON Schema validation · observability · latency & "
              "token-cost instrumentation")
    para(d, [EVIDENCE_BOUNDARY,
             ("portfolio repositories are independent work. Three years stated as three years. Agent-harness, CI, and "
              "verification depth is the claimed core; hyperscale distributed-systems depth (billions of calls) is NOT "
              "claimed — the harness and quality-gate engineering transfers, the scale is theirs to teach. No "
              "Vapi-internal claims. No Go / GCP / GraphQL / K8s / Terraform depth claimed.", False)])
    d.save(path)

if __name__ == "__main__":
    import os
    base = "/home/hatch/workspace/job-app-helix/packs-2026-10-21"
    build_pplx_behavior(f"{base}/perplexity-model-behavior/Casey_Barton_Resume_Perplexity_Model_Behavior.docx")
    build_pplx_evals(f"{base}/perplexity-answer-quality-evals/Casey_Barton_Resume_Perplexity_Answer_Quality_Evals.docx")
    build_vapi_dx(f"{base}/vapi-agentic-dx/Casey_Barton_Resume_Vapi_Agentic_DX.docx")
    print("built 3 docx")
