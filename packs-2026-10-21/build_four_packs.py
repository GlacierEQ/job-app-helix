"""Build four tailored resume packs from the 2026-10-21 04:53 HST discovery sweep.
Style replicates packs-2026-10-20/build_tonight_packs.py exactly.
Estate kit numbers (standing, live-verified 2026-10-20 ~23:57 HST): 65 MCP / 1398 repos / 989-test / 62-62.
Re-verified 2026-10-21 04:54-05:0x HST via JD page readbacks (all four reqs live).
Pack-building only — no submissions, ledger untouched.
"""
from docx import Document
from docx.shared import Pt, RGBColor

TEAL = RGBColor(0x0E, 0x5F, 0x6B)
GRAY = RGBColor(0x59, 0x59, 0x59)
INK = RGBColor(0x1A, 0x1A, 0x1A)
CONTACT = ("Honolulu, HI  ·  (808) 936-5654  ·  glacier.equilibrium@gmail.com  ·  "
           "github.com/GlacierEQ  ·  casey-barton-glaciereq.vercel.app  ·  linkedin.com/in/caseybartonai")

LS = 1.0708333333333333


def base_doc():
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.font.color.rgb = INK
    sec = d.sections[0]
    sec.left_margin = Pt(46.8); sec.right_margin = Pt(46.8)
    sec.top_margin = Pt(39.6); sec.bottom_margin = Pt(39.6)
    return d


def name_p(d, text):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = LS
    r = p.add_run(text); r.bold = True; r.font.size = Pt(25); r.font.color.rgb = TEAL
    return p


def contact_p(d):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = LS
    r = p.add_run(CONTACT); r.font.size = Pt(9.5); r.font.color.rgb = GRAY
    return p


def head(d, text):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(5); p.paragraph_format.line_spacing = LS
    r = p.add_run(text); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = TEAL
    return p


def para(d, runs, before=0, after=2):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = LS
    for text, bold in runs:
        r = p.add_run(text); r.bold = bold; r.font.size = Pt(10.5); r.font.color.rgb = INK
    return p


def comp_line(d, label, items):
    return para(d, [(label + ": ", True), (items, False)])


def bullet(d, runs):
    p = para(d, [("•  ", False)] + runs, after=2)
    p.paragraph_format.left_indent = Pt(0)
    return p


def exp_header(d, title, dates, org, loc):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = LS
    r1 = p.add_run(title + "\t"); r1.bold = True
    r2 = p.add_run(dates); r2.bold = True
    for r in (r1, r2):
        r.font.size = Pt(10.5); r.font.color.rgb = INK
    para(d, [(org + " — " + loc, False)], before=0, after=2)


def edu(d, rows):
    for school, dates, degree in rows:
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = LS
        r1 = p.add_run(school + "\t"); r1.bold = True; r1.font.size = Pt(10.5); r1.font.color.rgb = INK
        r2 = p.add_run(dates); r2.bold = True; r2.font.size = Pt(10.5); r2.font.color.rgb = INK
        para(d, [(degree, False)], before=0, after=2)


EDU_ROWS = [
    ("University of Hawaii at Manoa", "2016", "Bachelor of Science"),
    ("AWS Cloud Institute", "2025", "Cloud Computing program"),
    ("Colorado State University", "", "Undergraduate coursework"),
]

BOUNDARY_COMMON = ("portfolio repositories are independent work. Founder/operator depth (three years, 2023–present), "
                   "honestly stated — no enterprise customer logos, no enterprise-account track record invented. "
                   "No weight fine-tuning claimed — model strategy is eval-driven selection and routing. "
                   "No Go / GCP / GraphQL / K8s / Terraform / data-science-statistics depth claimed.")

PROFILE_COMMON_CLOSER = ("Honest on scope: no weight fine-tuning claimed — eval-driven selection and routing is the "
                         "proven part; founder/operator (three years, 2023–present), honestly stated.")


# ==================== 1. FLODESK — Senior Applied AI Engineer (Agent Quality & Evaluations) ====================
def build_flodesk(path):
    d = base_doc()
    name_p(d, "Casey Barton")
    contact_p(d)

    head(d, "PROFILE")
    para(d, [
        ("I own AI evaluation practice as a technical discipline, not a vibe: representative datasets, behavioral "
         "scenarios, scoring rubrics, regression suites, and human-review loops — with the eval harnesses, "
         "monitoring, and quality dashboards that make quality measurable without one-off manual checks. Three years "
         "as a founder running an evidence-grade agent control plane from Honolulu: I set quality baselines and "
         "fail-closed release gates for prompt, model, and agent changes, and I diagnose failures across prompts, "
         "context, orchestration, models, tools, and product code — then fix the right layer.", False),
    ])
    para(d, [
        ("The tuning is eval-driven, never claimed weight work: ", False), ("65", True),
        (" authored MCP connectors with typed contracts on every tool boundary, LLM-as-a-judge with written rubrics, "
         "golden datasets, and trace intelligence feeding prompt/retrieval/configuration iteration. ", False),
        (PROFILE_COMMON_CLOSER, False),
    ])

    head(d, "CORE COMPETENCIES")
    comp_line(d, "AI evaluation practice",
              "representative datasets · behavioral scenarios · scoring rubrics (LLM-as-a-judge) · regression suites · "
              "human review loops · quality baselines · release gates for prompt/model/agent changes · experimental design")
    comp_line(d, "Eval harnesses & observability",
              "eval harness design & operation · monitoring & quality dashboards · trace intelligence · latency/cost/"
              "quality instrumentation · failure analysis across prompts, context, orchestration, models, tools")
    comp_line(d, "Agentic systems",
              "agent loops & tools · structured outputs · multi-turn agents with tool use · user-intent preservation "
              "as behavioral assertions · MCP servers & connectors · guardrails · fail-closed recovery")
    comp_line(d, "Engineering",
              "Python (primary) · TypeScript · SQL · prompt/instruction versioning with change records · GitHub Actions "
              "(CI/CD) · Postgres · REST APIs · JSON Schema validation")

    head(d, "EXPERIENCE")
    exp_header(d, "Founder / Applied AI Systems Architect", "2023 – Present", "GlacierEQ", "Honolulu, HI")
    bullet(d, [("Own the evaluation practice end to end: define how the system should behave, write the quality "
                "definitions and scoring rubrics, and build the technical tooling — eval harnesses, regression "
                "suites, monitoring, quality dashboards — that make behavior measurable and the improvement loop "
                "continuous.", False)])
    bullet(d, [("Set quality baselines and release gates for prompt, model, and agent changes: every change carries a "
                "clear record of what changed, why it changed, and how behavior improved — fail-closed GO / NO-GO "
                "decisions when quality regresses.", False)])
    bullet(d, [("Diagnose failures across prompts, context, orchestration, models, tools, and product code: "
                "instrument the failure modes found in production — stuck loops, bad retrievals, tool misuse, "
                "context blowouts — and fix the right layer, with production failures and traces turned into durable "
                "regression cases.", False)])
    bullet(d, [("Combine automated checks with human judgment: ", False), ("989-test", True),
               (" suite with behavioral assertions, ", False), ("62/62", True),
               (" verification checks green, and human-review loops where judgment can't be automated — across an ",
                False), ("1398-repository", True), (" estate (live-verified 2026-10-20).", False)])
    bullet(d, [("Connect agents to tools and data through ", False), ("65", True),
               (" authored MCP connectors — governed connections with typed contracts, structured validation, and "
                "audit receipts on every boundary, so evaluation measures real tool behavior, not stubs.", False)])
    bullet(d, [("Tune with data, not opinion: model selection and routing across frontier + open models decided by "
                "measured quality, latency, and cost economics — eval-driven, never claimed weight fine-tuning.", False)])
    bullet(d, [("Translate ambiguous quality complaints into testable definitions: member-behavior ambiguity in, "
                "concrete system changes and regression coverage out — product sense expressed as written rubrics.",
                False)])

    head(d, "EDUCATION")
    edu(d, EDU_ROWS)

    head(d, "SELECTED PROOF")
    para(d, [("Evaluation-gated quality — measured before it ships", True)])
    bullet(d, [("Evaluation harnesses with written rubrics (LLM-as-a-judge) gate every deployment: golden datasets, "
                "behavioral assertions, fail-closed gates — the release-gate discipline applied to every run.", False)])
    para(d, [("Job-App Helix — evaluation-gated campaign engine", True)])
    bullet(d, [("Immutable evidence models with fail-closed GO / NO-GO decisions and per-run verification receipts — "
                "the continuous-improvement loop with the measurement built in.", False)])
    para(d, [("Failure posture — diagnose the right layer", True)])
    bullet(d, [("Production failure modes instrumented (stuck loops, bad retrievals, tool misuse, context blowouts) "
                "with replay loops and recovery restricted to evidence declared before reassessment — durable "
                "regression cases, not one-off fixes.", False)])

    head(d, "SKILLS")
    comp_line(d, "Evaluation", "golden datasets · behavioral scenarios · scoring rubrics · LLM-as-a-judge · regression suites · release gates · experimental design · human-review design")
    comp_line(d, "Agentic AI", "agent loops · tool/function calling · MCP · structured outputs · RAG · guardrails · fail-closed recovery · trace intelligence")
    comp_line(d, "Engineering", "Python · TypeScript · SQL · REST APIs · Postgres · CI/CD (GitHub Actions) · JSON Schema validation · observability · latency & token-cost instrumentation")
    para(d, [("Evidence boundary: ", True), (BOUNDARY_COMMON, False),
             (" Eval-platform experience is self-built harnesses with comparable mechanics — Langfuse/Braintrust "
              "listed as bonus only; no depth on those platforms claimed.", False)])
    d.save(path)


# ==================== 2. WHEEL — Staff Software Engineer, AI-Native Systems (Tech Lead) ====================
def build_wheel(path):
    d = base_doc()
    name_p(d, "Casey Barton")
    contact_p(d)

    head(d, "PROFILE")
    para(d, [
        ("Tech-lead posture from a founder's seat: I own AI-native domains outright — agent architecture, platform "
         "abstractions, evaluation and guardrail infrastructure — setting technical direction, writing the "
         "load-bearing design docs, and staying in the code for the hardest parts. Three years shipping production "
         "agentic systems from Honolulu: agents with retrieval, orchestration, policy-based routing, tool "
         "invocation, evaluation harnesses, and lifecycle observability — production-ready standards held in review: "
         "testability, rollback safety, cost ceilings, failure modes, human-in-the-loop boundaries.", False),
    ])
    para(d, [
        ("My platform layer is ", False), ("65", True),
        (" authored MCP connectors — governed tool/API/data-source connections with typed contracts, structured "
         "validation, and audit receipts — plus eval infrastructure (golden datasets, LLM-as-a-judge with written "
         "rubrics, regression gates) and reusable patterns every engagement inherits. ", False),
        (PROFILE_COMMON_CLOSER, False), (" No regulated-environment (HIPAA) production claimed — only the "
         "engineering posture transfers: data minimization, access controls, audit logging, and staged release "
         "discipline as architecture, not afterthoughts.", False),
    ])

    head(d, "CORE COMPETENCIES")
    comp_line(d, "Technical leadership",
              "domain ownership · tech-lead decomposition & sequencing · design docs · load-bearing architectural "
              "calls · ambiguity removal · cross-team influence without positional authority · mentorship by review")
    comp_line(d, "Agent architecture",
              "production AI agents · retrieval · orchestration · policy-based routing · tool invocation · "
              "evaluation harnesses · guardrail infrastructure · lifecycle observability · human-in-the-loop boundaries")
    comp_line(d, "AI platform foundations",
              "platform abstractions & shared libraries · 65 authored MCP connectors · typed contracts · reusable "
              "engineering patterns · privacy/security/compliance as architecture")
    comp_line(d, "Cloud-native & full-stack",
              "TypeScript/Node.js + Python · SQL (relational fluency) · REST APIs · CI/CD (GitHub Actions) · "
              "observability · on-call-ready operations · latency & token-cost instrumentation")

    head(d, "EXPERIENCE")
    exp_header(d, "Founder / Applied AI Systems Architect", "2023 – Present", "GlacierEQ", "Honolulu, HI")
    bullet(d, [("Own the technical direction for an AI-native domain end to end: agent architecture, platform "
                "abstractions, and evaluation/guardrail infrastructure across an ", False), ("1398-repository", True),
               (" estate (live-verified 2026-10-20) — decomposing ambiguous problems, sequencing delivery, clearing "
                "blockers, keeping the work pointed at the outcome rather than the ticket.", False)])
    bullet(d, [("Ship production AI agents incorporating retrieval, orchestration, policy-based routing, tool "
                "invocation, evaluation harnesses, and lifecycle observability — agents that carry real operational "
                "load, measured after the fact, not just at demo time.", False)])
    bullet(d, [("Set the production-ready bar and hold it: testability, rollback safety, cost ceilings, failure "
                "modes, human-in-the-loop boundaries — ", False), ("989-test", True), (" suite, ", False),
               ("62/62", True), (" verification checks green, deterministic CI regression gates, fail-closed "
                "recovery on every deployed agent.", False)])
    bullet(d, [("Build the abstraction layers teams build on: ", False), ("65", True),
               (" governed MCP connectors plus shared patterns, typed contracts, and structured validation — the "
                "reusable platform surface behind every deployed engagement.", False)])
    bullet(d, [("Own the evaluation strategy per domain: define the metrics, test harnesses, and evaluation plans "
                "that measure agent accuracy, latency, safety, and cost-effectiveness — kill approaches that don't "
                "work, early and visibly.", False)])
    bullet(d, [("Design for failure and advise as a technologist: instrument production failure modes — stuck loops, "
                "bad retrievals, tool misuse, context blowouts — with the catchers (replay loops, fail-closed "
                "gates, behavioral assertions), and run design sessions alongside operators and stakeholders so "
                "field friction becomes engineering requirements that move the roadmap.", False)])

    head(d, "EDUCATION")
    edu(d, EDU_ROWS)

    head(d, "SELECTED PROOF")
    para(d, [("Governed integration layer — 65 authored MCP connectors", True)])
    bullet(d, [("Typed contracts and structured validation on every tool boundary — the platform-abstraction "
                "discipline, including enterprise-style identity/access patterns: approval gates, access controls, "
                "audit receipts.", False)])
    para(d, [("Evaluation-gated quality — the measure-and-improve loop", True)])
    bullet(d, [("Golden datasets, behavioral assertions, LLM-as-a-judge with written rubrics, trace intelligence — "
                "the evaluation story behind every shipped agent; fail-closed GO / NO-GO decisions keep regression "
                "from carrying load.", False)])

    head(d, "SKILLS")
    comp_line(d, "Languages", "Python · TypeScript / Node.js · SQL · Bash · JavaScript")
    comp_line(d, "Agentic AI",
              "agent workflows · retrieval · orchestration · policy-based routing · tool/function calling · MCP · "
              "structured outputs · eval harnesses · guardrails · lifecycle observability")
    comp_line(d, "Platform & production",
              "platform abstractions · shared libraries · CI/CD (GitHub Actions) · Postgres · observability · "
              "retry logic · latency & token-cost instrumentation · design docs")
    para(d, [("Evidence boundary: ", True), (BOUNDARY_COMMON, False),
             (" Founder/operator depth is three years, 2023–present — stated against the JD's 8+ years "
              "(see REVIEW_REQUIRED.txt). No regulated/HIPAA production claimed; compliance-as-architecture is "
              "engineering practice, not domain experience.", False)])
    d.save(path)


# ==================== 3. FREDRIN — Forward Deployed Engineer ====================
def build_fredrin(path):
    d = base_doc()
    name_p(d, "Casey Barton")
    contact_p(d)

    head(d, "PROFILE")
    para(d, [
        ("Background-agent orchestration is my daily operating pattern: three years as a founder running many AI "
         "agents in parallel across a real codebase — real branches, worktrees, and review flow as the native "
         "medium. I own customer deployments end to end the founder way: make the repo agent-ready, wire the "
         "runtime layer, debug the whole seam in the field, and turn field learning into product changes I ship "
         "myself. Strong TypeScript and real git depth — branches, worktrees, rebases, and what breaks when many "
         "agents touch one repo.", False),
    ])
    para(d, [
        ("Close to the OS: processes, file systems, shells, and CLIs are where my agents live. Practical AI coding "
         "agents in real work — not demos — with eval harnesses and fail-closed gates keeping parallel work honest. "
         "Written-first communication, self-directed, the bias toward shipping something real every week. ", False),
        (PROFILE_COMMON_CLOSER, False),
    ])

    head(d, "CORE COMPETENCIES")
    comp_line(d, "Parallel-agent orchestration",
              "many agents in parallel · git worktrees & branches · rebase discipline · review flow · what breaks "
              "when many agents touch one repo · race/conflict detection · recovery")
    comp_line(d, "Forward-deployed delivery",
              "embedded delivery posture · customer discovery · sit-with-the-operator scoping · ambiguous-scope "
              "framing · rapid prototyping → production · handover discipline · field-to-product feedback")
    comp_line(d, "Engineering",
              "TypeScript · Python · SQL · Bash · processes/file systems/shells/CLIs · REST APIs · CI/CD "
              "(GitHub Actions) · JSON Schema validation · Postgres")
    comp_line(d, "Agent quality",
              "eval harnesses · golden datasets · behavioral assertions · LLM-as-a-judge with written rubrics · "
              "fail-closed recovery · trace intelligence · guardrails")

    head(d, "EXPERIENCE")
    exp_header(d, "Founder / Applied AI Systems Architect", "2023 – Present", "GlacierEQ", "Honolulu, HI")
    bullet(d, [("Run parallel-agent orchestration as the native development pattern: many agents working real "
                "branches and worktrees against a live codebase — the same mechanical discipline Fredrin productizes "
                "for customers (own toolchain, not Fredrin itself).", False)])
    bullet(d, [("Own the end-to-end deployment motion the founder way: make repos agent-ready, author the project "
                "context and conventions that stop workers guessing, wire the runtime layer (agents, connectors, "
                "automations against APIs), and stay until the work is merging.", False)])
    bullet(d, [("Debug the whole seam: shells, worktrees, branches, CI status — instrument what breaks when agents "
                "touch shared repos and build the catchers: race/conflict detection, replay loops, fail-closed "
                "gates.", False)])
    bullet(d, [("Turn field learning into product changes shipped directly: reusable AI-engineering playbooks, eval "
                "harnesses, and integration patterns — every new engagement inherits the stack, not a blank slate.",
                False)])
    bullet(d, [("Build the integration layer behind the agents: ", False), ("65", True),
               (" authored MCP connectors with typed contracts and structured validation — the connector discipline "
                "that turns recurring needs into reusable capabilities.", False)])
    bullet(d, [("Gate parallel work with evaluation: ", False), ("989-test", True), (" suite, ", False),
               ("62/62", True), (" verification checks green, behavioral assertions and written rubrics keeping "
                "agent output auditable — across an ", False), ("1398-repository", True), (" estate (live-verified "
                "2026-10-20).", False)])

    head(d, "EDUCATION")
    edu(d, EDU_ROWS)

    head(d, "SELECTED PROOF")
    para(d, [("Parallel-agent orchestration as native pattern", True)])
    bullet(d, [("Multi-agent loops with typed handoff contracts, structured outputs, and fail-closed recovery — "
                "agents working real branches and worktrees against live codebases, three years.", False)])
    para(d, [("Governed integration layer — 65 authored MCP connectors", True)])
    bullet(d, [("Governed tool/API/data-source connections with typed contracts — the runtime-layer wiring "
                "discipline, reusable across every engagement.", False)])
    para(d, [("Field-to-product loop", True)])
    bullet(d, [("Embedded operator scoping: sit with the workflow, instrument it, rebuild it — field friction "
                "captured and shipped as product changes, not left as one-off scripts.", False)])

    head(d, "SKILLS")
    comp_line(d, "Languages", "TypeScript · Python · SQL · Bash · JavaScript / Node.js")
    comp_line(d, "Git depth", "branches · worktrees · rebases · multi-agent repo conflict discipline · CI integration")
    comp_line(d, "Agent systems", "parallel agent orchestration · tool/function calling · MCP · structured outputs · eval harnesses · LLM-as-a-judge · fail-closed recovery")
    comp_line(d, "Production", "REST APIs · Postgres · CI/CD (GitHub Actions) · observability · processes & shells · JSON Schema validation")
    para(d, [("Evidence boundary: ", True), (BOUNDARY_COMMON, False),
             (" CRITICAL: no Fredrin usage claimed — the JD asks applicants to already use Fredrin; see "
              "REVIEW_REQUIRED.txt for the pre-application open call. Parallel-agent orchestration is his own "
              "toolchain, not Fredrin-internal experience.", False)])
    d.save(path)


# ==================== 4. STACKSYNC — Forward Deployed Engineer + Product Engineer ====================
def build_stacksync(path):
    d = base_doc()
    name_p(d, "Casey Barton")
    contact_p(d)

    head(d, "PROFILE")
    para(d, [
        ("Customer-facing builder with the full FDE motion: discovery, scope and success criteria, production "
         "integrations and AI workflows shipped against real data and unreliable APIs, and recurring customer needs "
         "turned into reusable connectors and platform capabilities — not one-offs left behind at accounts. Three "
         "years as a founder doing exactly this: write the code, make the technical decisions, stay accountable "
         "for the outcome, from Honolulu.", False),
    ])
    para(d, [
        ("My reusable-connector discipline is concrete: ", False), ("65", True),
        (" authored MCP connectors with typed contracts and structured validation — integrations built once and "
         "reused across every engagement. Quality is eval-gated before anything ships: golden datasets, "
         "behavioral assertions, LLM-as-a-judge with written rubrics, fail-closed recovery. ", False),
        (PROFILE_COMMON_CLOSER, False), (" No enterprise logos claimed; demos and POCs are founder-scale, honestly "
        "stated.", False),
    ])

    head(d, "CORE COMPETENCIES")
    comp_line(d, "Customer implementation",
              "technical discovery · scope & success-criteria framing · customized demos from real workflows · "
              "production integrations · deployment & iteration · operator/executive communication")
    comp_line(d, "Reusable platform building",
              "recurring needs → reusable connectors · 65 authored MCP connectors · typed contracts · platform "
              "capabilities & developer experience · field-to-product feedback loops")
    comp_line(d, "Agentic AI",
              "AI agents on real data · workflows & automation · tool use & function calling · structured outputs · "
              "MCP · guardrails · fail-closed recovery · human-in-the-loop gates")
    comp_line(d, "Engineering",
              "Python (primary) · TypeScript · SQL · REST APIs · Postgres · databases & integrations · "
              "CI/CD (GitHub Actions) · observability · latency & token-cost instrumentation · retry logic")

    head(d, "EXPERIENCE")
    exp_header(d, "Founder / Applied AI Systems Architect", "2023 – Present", "GlacierEQ", "Honolulu, HI")
    bullet(d, [("Own customer implementations end to end: lead discovery, define scope and success criteria, design "
                "the solution, take it through deployment and iteration — the full path from 'here's how the "
                "business works' to a live, reliable implementation.", False)])
    bullet(d, [("Ship production systems against real data: integrations, workflows, and AI agents handling "
                "unreliable APIs, edge cases, and operational constraints — measured before and after, not vibes.",
                False)])
    bullet(d, [("Build the product, not just one-off solutions: turn recurring customer needs into reusable "
                "connectors, platform capabilities, and better developer experience — ", False), ("65", True),
               (" governed MCP connectors as the standing proof of that discipline.", False)])
    bullet(d, [("Build demos that make the value concrete: prospect workflows turned into working demonstrations "
                "with declared success criteria — field friction instrumented and written into engineering-ready "
                "requirements.", False)])
    bullet(d, [("Gate quality with evaluation harnesses: golden datasets, behavioral assertions, LLM-as-a-judge with "
                "written rubrics — ", False), ("989-test", True), (" suite, ", False), ("62/62", True),
               (" verification checks green across an ", False), ("1398-repository", True), (" estate (live-verified "
                "2026-10-20).", False)])
    bullet(d, [("Write the code, make the technical decisions, stay accountable for the outcome: founder/operator "
                "delivery, three years — ambiguity in, working software out, without waiting for detailed "
                "instructions.", False)])

    head(d, "EDUCATION")
    edu(d, EDU_ROWS)

    head(d, "SELECTED PROOF")
    para(d, [("Governed integration layer — 65 authored MCP connectors", True)])
    bullet(d, [("Governed tool/API/data-source connections with typed contracts — the reusable-connector discipline "
                "the role asks for, proven across every engagement.", False)])
    para(d, [("Evaluation-gated quality — shipped measured", True)])
    bullet(d, [("Eval harnesses with written rubrics gate every deployment: judged criteria declared up front, "
                "behavioral assertions enforced, trace intelligence feeding iteration.", False)])
    para(d, [("Customer-first delivery", True)])
    bullet(d, [("Three years of founder/operator embedded delivery: discovery, customized demos, solution design, "
                "deployment and iteration — customer-first with good questions and clear communication.", False)])

    head(d, "SKILLS")
    comp_line(d, "Languages", "Python · TypeScript · SQL · Bash · JavaScript / Node.js")
    comp_line(d, "Integrations", "65 authored MCP connectors · REST APIs · CRM/database integration patterns · real-time sync posture · structured validation")
    comp_line(d, "Agentic AI", "AI agents · workflows & automation · tool/function calling · MCP · structured outputs · eval harnesses · fail-closed recovery")
    comp_line(d, "Production", "Postgres · CI/CD (GitHub Actions) · observability · latency & token-cost instrumentation · retry logic · technical writing")
    para(d, [("Evidence boundary: ", True), (BOUNDARY_COMMON, False),
             (" No enterprise customer logos or enterprise-account track record claimed. Travel (US + SF office) "
              "is a JD core requirement — see REVIEW_REQUIRED.txt; Honolulu-based, his call.", False)])
    d.save(path)


BASE = "/home/hatch/workspace/job-app-helix/packs-2026-10-21"

if __name__ == "__main__":
    import os
    jobs = [
        ("flodesk-senior-applied-ai-eval", build_flodesk, "Flodesk_Senior_Applied_AI_Eval"),
        ("wheel-staff-ai-native-systems", build_wheel, "Wheel_Staff_AI_Native_Systems"),
        ("fredrin-forward-deployed-engineer", build_fredrin, "Fredrin_Forward_Deployed_Engineer"),
        ("stacksync-fde", build_stacksync, "Stacksync_FDE"),
    ]
    for slug, fn, stem in jobs:
        os.makedirs(f"{BASE}/{slug}", exist_ok=True)
        fn(f"{BASE}/{slug}/Casey_Barton_Resume_{stem}.docx")
    print("built 4 docx")
