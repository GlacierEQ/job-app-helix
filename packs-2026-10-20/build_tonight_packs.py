"""Build the two DRAFT resume packs for 2026-10-20 23:5x sweep shortlist.
Style replicates packs-2026-10-20/cohere-fde-agentic/ exactly.
All estate numbers live-verified 2026-10-20 ~23:57 HST (gh GraphQL viewer.repositories.totalCount:
1,398 repos; kit-allowlist: 65 MCP / 989-test / 62-62). Estate repo figure uses LIVE 1,398
(corrected 23:58 — an earlier draft used 886 from a different query; re-verified direct readback wins).
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

TEAL = RGBColor(0x0E, 0x5F, 0x6B)
GRAY = RGBColor(0x59, 0x59, 0x59)
INK = RGBColor(0x1A, 0x1A, 0x1A)
CONTACT = ("Honolulu, HI  ·  (808) 936-5654  ·  glacier.equilibrium@gmail.com  ·  "
           "github.com/GlacierEQ  ·  casey-barton-glaciereq.vercel.app  ·  linkedin.com/in/caseybartonai")

LS = 1.0708333333333333


def base_doc():
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.font.color.rgb = INK
    sec = d.sections[0]
    sec.left_margin = Pt(46.8); sec.right_margin = Pt(46.8)
    sec.top_margin = Pt(39.6); sec.bottom_margin = Pt(39.6)
    return d


def name_p(d, text):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = LS
    r = p.add_run(text); r.bold = True; r.font.size = Pt(25); r.font.color.rgb = TEAL
    return p


def contact_p(d):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = LS
    r = p.add_run(CONTACT); r.font.size = Pt(9.5); r.font.color.rgb = GRAY
    return p


def head(d, text):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(5); p.paragraph_format.line_spacing = LS
    r = p.add_run(text); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = TEAL
    return p


def para(d, runs, before=0, after=2, bold_first=False):
    """runs: list of (text, bold)."""
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = LS
    for text, bold in runs:
        r = p.add_run(text); r.bold = bold; r.font.size = Pt(10.5); r.font.color.rgb = INK
    return p


def comp_line(d, label, items):
    return para(d, [(label + ": ", True), (items, False)])


def bullet(d, runs):
    p = para(d, [("•  ", False)] + runs, after=2)
    p.paragraph_format.left_indent = Pt(0)
    return p


def exp_header(d, title, dates, org, loc):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = LS
    r1 = p.add_run(title + "\t"); r1.bold = True
    r2 = p.add_run(dates); r2.bold = True
    for r in (r1, r2):
        r.font.size = Pt(10.5); r.font.color.rgb = INK
    para(d, [(org + " — " + loc, False)], before=0, after=2)


def edu(d, rows):
    for school, dates, degree in rows:
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = LS
        r1 = p.add_run(school + "\t"); r1.bold = True; r1.font.size = Pt(10.5); r1.font.color.rgb = INK
        r2 = p.add_run(dates); r2.bold = True; r2.font.size = Pt(10.5); r2.font.color.rgb = INK
        para(d, [(degree, False)], before=0, after=2)


# ==================== ATOMIC VC ====================
def build_atomic(path):
    d = base_doc()
    name_p(d, "Casey Barton")
    contact_p(d)

    head(d, "PROFILE")
    para(d, [
        ("Builder-first founding AI engineer. Three years running an evidence-grade agent control plane from "
         "Honolulu — I ship domain-specific agents into live operator workflows, then make them repeatable: "
         "evaluation infrastructure that measures whether agents actually replace human workflows (not vibes), "
         "model selection and routing across frontier + open models decided by measured economics, and AI-engineering "
         "playbooks every new engagement inherits on day one. 90% build is my natural ratio. Field fluency runs "
         "deeper than theory: I come from the physical economy — licensed home-inspector background, permits, "
         "site work, materials — so unglamorous industries are home turf.", False),
    ])
    para(d, [
        ("My integration layer is ", False), ("65", True),
        (" authored MCP connectors — governed tool/API/data-source connections with typed contracts and structured "
         "validation on every boundary — plus eval harnesses with written rubrics (LLM-as-a-judge), golden datasets, "
         "and behavioral assertions. Agents must be reliable, observable, and auditable from day one: the operating "
         "posture, not a compliance add-on. Honest on scope: no client-side enterprise deployments, no enterprise "
         "logos claimed; the building, the evals, and the playbooks are the transferable part.", False),
    ])

    head(d, "CORE COMPETENCIES")
    comp_line(d, "Forward-deployed delivery",
              "embedded delivery posture · customer discovery · sit-with-the-customer workflow instrumentation · "
              "ambiguous-scope scoping · rapid prototyping → production · legacy workflow replacement · handover discipline")
    comp_line(d, "Agentic AI",
              "domain-specific agents (reason / plan / act) · tool use & function calling · structured outputs · "
              "MCP servers & connectors · RAG pipelines · multi-agent loops · guardrails · fail-closed recovery · human-in-the-loop gates")
    comp_line(d, "Evaluation, quality & model economics",
              "eval infrastructure · golden datasets · behavioral assertions · LLM-as-a-judge with written rubrics · "
              "success-criteria framing (agents replacing workflows, measured) · frontier + open routing by "
              "cost/quality economics · trace intelligence")
    comp_line(d, "Production",
              "Python (primary) · TypeScript · SQL · Bash · REST APIs · JSON Schema / Pydantic-style validation · "
              "GitHub Actions (CI/CD) · Postgres · per-task latency & token-cost instrumentation · retry logic")

    head(d, "EXPERIENCE")
    exp_header(d, "Founder / Applied AI Systems Architect", "2023 – Present", "GlacierEQ", "Honolulu, HI")
    bullet(d, [("Ship working agents into live environments fast: the founding motion is domain-specific agent "
                "prototyped with the operator, measured against their real workflow, then hardened — "
                "30-day-style delivery cadence across engagements.", False)])
    bullet(d, [("Build the eval infrastructure that measures whether agents replace human workflows — not vibes: "
                "golden datasets, behavioral assertions, LLM-as-a-judge with written rubrics, and trace intelligence "
                "feeding prompt/retrieval/configuration iteration.", False)])
    bullet(d, [("Decide fine-tune vs. prompt vs. RAG vs. agent vs. harness with numbers: model strategy is eval-driven "
                "selection and routing across frontier + open models by measured quality, latency, and cost — "
                "opinions defended by the data, not the trend. (No weight fine-tuning claimed; the routing and "
                "selection economics are the proven part.)", False)])
    bullet(d, [("Sit with the operator, watch the workflow, instrument it, rebuild it: embedded scoping turns "
                "ambiguous field problems into well-framed agent deployments with clear success criteria.", False)])
    bullet(d, [("Make it repeatable: reusable AI-engineering playbooks, model stacks, eval harnesses, and integration "
                "patterns — every new engagement inherits the stack and the lessons, not a blank slate.", False)])
    bullet(d, [("Connect agents to the real world through ", False), ("65", True),
               (" authored MCP connectors — governed integrations with typed contracts, evidence boundaries, and "
                "audit receipts on every tool boundary.", False)])
    bullet(d, [("Run an evidence-grade agent control plane across an ", False), ("1398-repository", True),
               (" estate (live-verified 2026-10-20): observability, trace intelligence, recoverable runs — backed by "
                "a ", False), ("989-test", True), (" suite, ", False), ("62/62", True),
               (" verification checks green, and deterministic CI regression gates.", False)])

    head(d, "EDUCATION")
    edu(d, [
        ("University of Hawaii at Manoa", "2016", "Bachelor of Science"),
        ("AWS Cloud Institute", "2025", "Cloud Computing program"),
        ("Colorado State University", "", "Undergraduate coursework"),
    ])

    head(d, "SELECTED PROOF")
    para(d, [("Job-App Helix — evaluation-gated campaign engine", True)])
    bullet(d, [("Attributable campaign engine with immutable evidence models and fail-closed GO / NO-GO decisions — "
                "the eval gate decides what carries load; written rubrics behind every automated quality decision.", False)])
    para(d, [("Governed integration layer — 65 authored MCP connectors", True)])
    bullet(d, [("Governed tool/API/data-source connections with typed contracts and structured validation on every "
                "boundary — the reusable integration layer behind every deployed agent engagement.", False)])
    para(d, [("Failure posture — ", True), ("designed for production, not the demo: instrument the failure modes found "
                "in production — stuck loops, bad retrievals, tool misuse, context blowouts — and build the catchers: "
                "replay loops, fail-closed gates, recovery restricted to evidence declared before reassessment.", False)])

    head(d, "SKILLS")
    comp_line(d, "Languages", "Python · TypeScript · SQL · Bash · JavaScript / Node.js")
    comp_line(d, "Agentic AI",
              "domain-specific agents · tool/function calling · MCP · structured outputs · RAG · eval harnesses · "
              "LLM-as-a-judge · trace intelligence · guardrails · fail-closed recovery")
    comp_line(d, "Deployed engineering",
              "discovery & scoping · workflow instrumentation · rapid prototyping · embedded delivery · "
              "field-to-product feedback · handover discipline · technical writing")
    comp_line(d, "Production",
              "REST APIs · Postgres · CI/CD (GitHub Actions) · JSON Schema validation · observability · "
              "latency & token-cost instrumentation · retry logic")
    para(d, [("Evidence boundary: ", True),
             ("portfolio repositories are independent work. No enterprise customer logos claimed; client-side deployment "
              "depth is founder/operator (three years), honestly stated. No weight fine-tuning claimed — model strategy "
              "is eval-driven selection and routing. Home-inspector background = construction inspection — physical-economy "
              "fluency, not engineering depth. No Go / GCP / GraphQL / K8s / Terraform / DS depth claimed.", False)])
    d.save(path)


# ==================== OPENROUTER ====================
def build_openrouter(path):
    d = base_doc()
    name_p(d, "Casey Barton")
    contact_p(d)

    head(d, "PROFILE")
    para(d, [
        ("Technical quarterback for AI sales cycles: I translate a prospect's architecture, use cases, and constraints "
         "into demos that answer their actual questions, solution designs they can defend internally, and POCs scoped "
         "with honest success criteria. Three years as a founder doing exactly that motion — technical discovery, "
         "customized demonstrations, architecture design under pressure, and technical objections handled with measured "
         "evidence instead of slides. The feedback loop runs both ways: field friction becomes product requirements "
         "written for engineers.", False),
    ])
    para(d, [
        ("The hands are real: ", False), ("65", True),
        (" authored MCP connectors — governed API/data-source integrations with typed contracts, structured "
        "validation, auth/access patterns, and audit receipts — plus integration-side work with the OpenRouter API "
        "(estate repo, independent work). Eval harnesses with written rubrics (LLM-as-a-judge) keep every demo and POC "
        "honest: measured latency, cost, and quality, not claims. Customer-first founder/operator depth (three years), "
        "honestly stated — no enterprise sales-engineering track record, no enterprise logos claimed.", False),
    ])

    head(d, "CORE COMPETENCIES")
    comp_line(d, "Customer engineering",
              "technical discovery · solution architecture · customized demos · POC environments · technical objections "
              "· competitive positioning · executive-friendly explanation · field-to-product feedback")
    comp_line(d, "API & integration",
              "65 authored MCP connectors · REST APIs · typed contracts · structured validation · auth/access patterns · "
              "SLA-shaped thinking · multi-provider routing patterns (integration-side)")
    comp_line(d, "Agentic AI",
              "agentic workflows (reason / plan / act) · tool use & function calling · structured outputs · MCP · "
              "RAG · guardrails · fail-closed recovery · human-in-the-loop gates")
    comp_line(d, "Evaluation & quality",
              "eval harnesses · golden datasets · behavioral assertions · LLM-as-a-judge with written rubrics · "
              "success-criteria framing · latency & token-cost instrumentation · failure analysis")
    comp_line(d, "Production",
              "Python (primary) · TypeScript · SQL · Bash · JSON Schema / Pydantic-style validation · GitHub Actions "
              "(CI/CD) · Postgres · retry logic · observability")

    head(d, "EXPERIENCE")
    exp_header(d, "Founder / Applied AI Systems Architect", "2023 – Present", "GlacierEQ", "Honolulu, HI")
    bullet(d, [("Lead the customer-engineering motion end to end: technical discovery with operators, customized "
                "demonstrations built around their actual use cases, solution architectures they can defend "
                "internally, and POC environments scoped with written success criteria.", False)])
    bullet(d, [("Answer technical objections with measured evidence, not decks: latency, cost, and quality numbers "
                "from real runs — routing, model-selection, and architecture tradeoffs defended by the data.", False)])
    bullet(d, [("Integrate deeply with customer stacks through ", False), ("65", True),
               (" authored MCP connectors — governed API/data-source connections with typed contracts, structured "
                "validation, and audit receipts on every boundary; enterprise-style identity/access patterns "
                "(approval gates, access controls) are native to the integration layer.", False)])
    bullet(d, [("Build integration-side familiarity with model-routing platforms: estate work exercises provider-routing "
                "APIs (OpenRouter integration-side, independent work) — model switching by cost/quality economics, "
                "failover and latency instrumentation.", False)])
    bullet(d, [("Close the field-to-product loop: operator friction captured, instrumented, and written into "
                "engineering-ready requirements — customer feedback that actually changes what gets built.", False)])
    bullet(d, [("Gate every shipped demo and deployment with evaluation harnesses: golden datasets, behavioral "
                "assertions, LLM-as-a-judge with written rubrics — ", False), ("989-test", True), (" suite, ", False),
               ("62/62", True), (" verification checks green, fail-closed recovery across an ", False),
               ("1398-repository", True), (" estate (live-verified 2026-10-20).", False)])

    head(d, "EDUCATION")
    edu(d, [
        ("University of Hawaii at Manoa", "2016", "Bachelor of Science"),
        ("AWS Cloud Institute", "2025", "Cloud Computing program"),
        ("Colorado State University", "", "Undergraduate coursework"),
    ])

    head(d, "SELECTED PROOF")
    para(d, [("Governed integration layer — 65 authored MCP connectors", True)])
    bullet(d, [("Governed tool/API/data-source connections with typed contracts and structured validation on every "
                "boundary — the reusable integration layer behind every deployed agent engagement.", False)])
    bullet(d, [("Connectors span enterprise-style identity/access patterns: approval gates, access controls, audit "
                "receipts — the auth and governance surface customer engineers actually demo.", False)])
    para(d, [("Evaluation-gated quality — demos measured, not claimed", True)])
    bullet(d, [("Evaluation harnesses with written rubrics gate every demo and deployment: judge criteria declared up "
                "front, behavioral assertions enforced, trace intelligence feeding iteration.", False)])
    bullet(d, [("Job-App Helix campaign engine — immutable evidence models, fail-closed GO / NO-GO decisions: the "
                "POC discipline of declared success criteria, applied to every run.", False)])

    head(d, "SKILLS")
    comp_line(d, "Languages", "Python · TypeScript · SQL · Bash · JavaScript / Node.js")
    comp_line(d, "Customer engineering",
              "technical discovery · demos & POCs · solution architecture · technical objections · competitive "
              "positioning · customer success posture · technical writing")
    comp_line(d, "Agentic AI",
              "agent workflows · tool/function calling · MCP · structured outputs · RAG · eval harnesses · "
              "LLM-as-a-judge · trace intelligence · guardrails")
    comp_line(d, "Production",
              "REST APIs · Postgres · CI/CD (GitHub Actions) · JSON Schema validation · observability · "
              "latency & token-cost instrumentation · retry logic")
    para(d, [("Evidence boundary: ", True),
             ("portfolio repositories are independent work. Founder/operator customer-facing depth (three years), "
              "honestly stated — no enterprise sales-engineering track record, no enterprise customer logos claimed. "
              "OpenRouter experience is integration-side (as a builder on the API), not OpenRouter-internal platform "
              "work. No Go / GCP / GraphQL / K8s / Terraform / data-science-statistics depth claimed. GitHub is the "
              "forge of record.", False)])
    d.save(path)


if __name__ == "__main__":
    build_atomic("/home/hatch/workspace/job-app-helix/packs-2026-10-20/atomic-vc-founding-fde/Casey_Barton_Resume_Atomic_Founding_FDE.docx")
    build_openrouter("/home/hatch/workspace/job-app-helix/packs-2026-10-20/openrouter-customer-engineer/Casey_Barton_Resume_OpenRouter_Customer_Engineer.docx")
    print("built both docx")
